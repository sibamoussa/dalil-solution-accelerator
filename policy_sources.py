"""Bounded public-site discovery and extraction; never carries Azure credentials."""

import hashlib
import io
import ipaddress
import json
import socket
import sys
import zipfile
from collections import deque
from pathlib import Path
from urllib.parse import unquote, urldefrag, urljoin, urlsplit
from urllib.robotparser import RobotFileParser

import certifi
import urllib3
from bs4 import BeautifulSoup
from docx import Document

sys.path.insert(0, str(Path(__file__).parent / "webapp"))
from accelerator_config import load_config


MAX_BYTES = 50 * 1024 * 1024
USER_AGENT = "DalilPolicyImporter/1.0"


def validate_url(url: str, source: dict, check_path: bool = True) -> str:
    parsed = urlsplit(url)
    if parsed.scheme != "https" or parsed.hostname not in source["allowedHosts"] or parsed.username or parsed.password or parsed.port not in (None, 443):
        raise ValueError("Source URL must use HTTPS on an explicitly allowed host")
    path = unquote(parsed.path)
    if "\\" in path or any(part in (".", "..") for part in path.split("/")):
        raise ValueError("Ambiguous source URL path")
    if check_path and not any(path.startswith(prefix) for prefix in source["pathPrefixes"]):
        raise ValueError("Source URL is outside configured path prefixes")
    return urldefrag(url)[0]


def public_addresses(host: str) -> list[str]:
    addresses = sorted({entry[4][0] for entry in socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)})
    if not addresses or any(not ipaddress.ip_address(address).is_global for address in addresses):
        raise ValueError("Refusing a private, loopback, reserved or mixed DNS destination")
    return addresses


class PublicWebsite:
    def __init__(self, source: dict):
        self.source = source
        self.robots = {}

    def fetch(self, url: str, robots: bool = True) -> tuple[bytes, str, str]:
        for attempt in range(5):
            url = validate_url(url, self.source, check_path=robots)
            parsed = urlsplit(url)
            if robots and not self.allowed_by_robots(url):
                raise ValueError("Source robots.txt disallows this document")
            address = public_addresses(parsed.hostname)[0]
            pool = urllib3.HTTPSConnectionPool(
                address, port=443, server_hostname=parsed.hostname, assert_hostname=parsed.hostname,
                cert_reqs="CERT_REQUIRED", ca_certs=certifi.where(), retries=False,
                timeout=urllib3.Timeout(connect=15, read=90),
            )
            response = None
            try:
                target = parsed.path or "/"
                if parsed.query:
                    target += "?" + parsed.query
                response = pool.request("GET", target, headers={"Host": parsed.hostname, "User-Agent": USER_AGENT, "Accept-Encoding": "identity"}, redirect=False, preload_content=False)
                if response.status in (301, 302, 303, 307, 308):
                    url = urljoin(url, response.headers["Location"])
                    continue
                if not robots and response.status == 404:
                    return b"", "text/plain", url
                if response.status != 200:
                    raise ValueError(f"Source request failed (HTTP {response.status})")
                limit = MAX_BYTES if robots else 1024 * 1024
                if int(response.headers.get("Content-Length", "0")) > limit:
                    raise ValueError("Source exceeds download limit")
                payload = bytearray()
                for block in response.stream(64 * 1024, decode_content=True):
                    payload.extend(block)
                    if len(payload) > limit:
                        raise ValueError("Source exceeds download limit")
                return bytes(payload), response.headers.get("Content-Type", "").split(";")[0].lower(), url
            finally:
                if response:
                    response.close()
                pool.close()
        raise ValueError("Too many source redirects")

    def allowed_by_robots(self, url: str) -> bool:
        parsed = urlsplit(url)
        origin = f"https://{parsed.netloc}"
        if origin not in self.robots:
            payload, _, _ = self.fetch(origin + "/robots.txt", robots=False)
            parser = RobotFileParser()
            parser.parse(payload.decode("utf-8", errors="replace").splitlines())
            self.robots[origin] = parser
        return self.robots[origin].can_fetch(USER_AGENT, url)


def extract(payload: bytes, mime: str, url: str) -> tuple[str, bytes, str]:
    if payload.startswith(b"%PDF-"):
        return "pdf", payload, Path(urlsplit(url).path).stem
    if payload.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > MAX_BYTES or "word/document.xml" not in archive.namelist():
                raise ValueError("Unsupported or oversized DOCX archive")
        document = Document(io.BytesIO(payload))
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        text = "\n".join(parts).strip()
        if not text:
            raise ValueError("DOCX contains no extractable text")
        return "docx", text.encode("utf-8"), Path(urlsplit(url).path).stem
    if mime not in ("text/html", "application/xhtml+xml"):
        raise ValueError("Only PDF, DOCX and HTML sources are supported")
    soup = BeautifulSoup(payload, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else Path(urlsplit(url).path).name
    for element in soup.select("script,style,nav,header,footer,aside,form,noscript"):
        element.decompose()
    content = soup.find("main") or soup.find("article") or soup.body or soup
    text = content.get_text("\n", strip=True)
    return "html", text.encode("utf-8"), title[:300]


def discover(config: dict, website: PublicWebsite | None = None) -> dict:
    source = config["source"]
    website = website or PublicWebsite(source)
    queue = deque([source["website"]])
    visited = set()
    documents = []
    skipped = []
    html_pages = 0
    while queue and len(documents) < source["maxDocuments"] and len(visited) < source["maxDocuments"] + source["maxPages"]:
        url = queue.popleft()
        if url in visited:
            continue
        visited.add(url)
        try:
            payload, mime, final_url = website.fetch(url)
            kind, content, title = extract(payload, mime, final_url)
            if kind == "html":
                if html_pages >= source["maxPages"]:
                    continue
                html_pages += 1
                soup = BeautifulSoup(payload, "html.parser")
                for link in soup.select("a[href]"):
                    candidate = urldefrag(urljoin(final_url, link["href"]))[0]
                    try:
                        validate_url(candidate, source)
                    except ValueError:
                        continue
                    if candidate not in visited and candidate not in queue and len(queue) < 2000:
                        queue.append(candidate)
                if len(content.decode("utf-8")) < 200:
                    continue
            filename = "policy-" + hashlib.sha256(final_url.encode()).hexdigest()[:24] + (".pdf" if kind == "pdf" else ".txt")
            if any(item["fileName"] == filename for item in documents):
                continue
            documents.append({"url": final_url, "title": title or final_url, "fileName": filename, "format": kind, "language": "", "sha256": hashlib.sha256(payload).hexdigest()})
        except (ValueError, OSError, urllib3.exceptions.HTTPError, zipfile.BadZipFile) as error:
            skipped.append({"url": url, "reason": str(error)})
    if not documents:
        raise ValueError("No supported documents found. Review path prefixes, robots.txt and website accessibility.")
    return {"schemaVersion": 1, "source_page": source["website"], "documents": documents, "skipped": skipped, "truncated": bool(queue)}


def main():
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = discover(load_config(args.config))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Discovered {len(result['documents'])} documents; {len(result['skipped'])} skipped; truncated={result['truncated']}. Review {args.output} before ingestion.")


if __name__ == "__main__":
    main()